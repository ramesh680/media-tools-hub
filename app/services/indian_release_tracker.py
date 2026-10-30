"""Indian Regional Release Tracker - web version of the `indian-regional-release-tracker` skill.

Builds one table of Indian movies and TV shows across regional languages for a chosen window,
or enriches only an uploaded title list, with the same columns as the skill's workbook:
IMDb ID, release date/status, genre, language, producer, Studio, Network, director, cast,
trailer link, Top 3 YouTube channels (``<channel URL>|<Title>``), CBFC certificate fields,
origin, Pay 1/Library window and source.

Run modes (same as the skill's menu):
  1  2024-01-01 -> today + 1 year (announced "Date TBA" titles kept)
  2  1st of this month -> today + 3 months
  3  1st of this month -> today + 6 months
  4  only the uploaded / pasted title list (any year, any country)
  5  custom: own date range, languages, TV on/off, notes

Server-side sources (all public, no login):
  * Wikipedia "List of <Language> films of <Year>" tables (MediaWiki API, CC BY-SA)
  * Wikipedia OTT-originals lists for TV shows (Netflix India, Prime Video India, JioHotstar)
  * IMDb Advanced Title Search (IMDb's public GraphQL, the API behind imdb.com/search/title):
    every title with India as a country of origin released in the window - movies + TV movies,
    and TV series + mini-series - with director, cast, producers, production companies (Studio),
    streaming/TV distributors (Network), runtime and certificate. Best effort: IMDb may refuse
    datacentre traffic; the run then notes it and continues
  * TMDB (TMDB_API_KEY / TMDB_READ_ACCESS_TOKEN) and OMDb (OMDB_API_KEY) - the hub's existing keys -
    fill missing IMDb IDs, genres, directors, cast, production companies, TV networks and the
    India streaming provider, so rows carry what the BDR ingest template needs
  * Wikidata SPARQL (CC0) to resolve uploaded titles <-> IMDb IDs, studio, network
  * YouTube search results page (no key) for the trailer + Top 3 channels; YouTube Data API v3
    (YOUTUBE_API_KEY or the per-run key) as a fallback when the results page is blocked
  * CBFC official website "Search Film" - best effort, read live on every run (no PDFs)

No network access happens at import time; every fetch goes through the shared HttpClient.
"""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO, StringIO
from typing import Any, Callable
import csv
import json
import re
import time
from urllib.parse import quote, quote_plus

from bs4 import BeautifulSoup
from dateutil import parser as dateparser
from dateutil.relativedelta import relativedelta

from app.models import utc_now_iso

ProgressCallback = Callable[[int, str], None]

TRACKER_TYPE = "indian_release_tracker"

INDIAN_LANGUAGES = ["Hindi", "Tamil", "Telugu", "Malayalam", "Kannada", "Marathi", "Bengali",
                    "Gujarati", "Punjabi", "Odia", "Bhojpuri", "Assamese", "Tulu"]
INDIAN_LOWER = {x.lower() for x in INDIAN_LANGUAGES} | {"urdu", "konkani", "haryanvi", "rajasthani", "sanskrit"}

TITLE_COLUMNS = [
    "Title Name", "IMDb ID", "Category", "Release Date", "Release Status", "Genre", "Language",
    "Produce by", "Studio", "Network", "Director", "Cast", "Teaser / Trailer on YouTube",
    "Top 3 YouTube Channels", "Type / Platform", "Certificate Type", "Certificate No.", "Duration",
    "CBFC Office", "Origin", "OTT Window (Pay 1 / Library)", "Source",
]
COVERAGE_COLUMNS = ["Title", "IMDb ID", "Language", "Year", "Origin", "Window", "Wikidata",
                    "Sources found", "Notes"]

STATUSES = ("Released", "Certified (CBFC)", "Confirmed", "Tentative", "Date TBA")
NETWORK_RE = re.compile(
    r"netflix|amazon|prime video|hotstar|jiocinema|disney|zee5|zee |sonyliv|sony |sun (tv|nxt)|\baha\b|"
    r"hoichoi|etv|star |colors|manorama|chaupal|planet marathi|mx player|apple tv|lionsgate play|"
    r"discovery|asianet|gemini|surya tv|udaya|waves ott", re.I)
MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

WIKI_API = "https://en.wikipedia.org/w/api.php"
WIKIDATA_SPARQL = "https://query.wikidata.org/sparql"
IMDB_GRAPHQL = "https://caching.graphql.imdb.com/"
YT_RESULTS = "https://www.youtube.com/results?search_query="
TMDB_API = "https://api.themoviedb.org/3"
OMDB_API = "https://www.omdbapi.com/"
YT_API_SEARCH = "https://www.googleapis.com/youtube/v3/search"
YT_API_CHANNELS = "https://www.googleapis.com/youtube/v3/channels"
CBFC_HOME = "https://www.cbfcindia.gov.in/"
CBFC_SEARCH_CANDIDATES = [
    "https://www.cbfcindia.gov.in/cbfcAdmin/search-film.php",
    "https://www.cbfcindia.gov.in/cbfcAdmin/search-result.php",
]
TV_LISTS = [
    ("List of Netflix India original programming", "Netflix"),
    ("List of Amazon India originals", "Prime Video"),
    ("List of Disney+ Hotstar original programming", "JioHotstar"),
]

# Free-tier friendly ceilings (a full mode-1 run otherwise means ~12k YouTube requests).
YOUTUBE_WORKERS = 4
YOUTUBE_MAX_TITLES = 600
YOUTUBE_API_MAX_TITLES = 25          # 3 searches x 100 quota units each -> ~7.5k units
CBFC_MAX_TITLES = 150
IMDB_SEARCH_PAGE = 250
IMDB_SEARCH_MAX_PAGES = 24            # 6,000 titles per kind is plenty for India in any window
TMDB_MAX_TITLES = 900
TMDB_WORKERS = 6
OMDB_MAX_TITLES = 150                 # free OMDb key: 1,000 requests a day

LANG_CODES = {"Hindi": "hi", "Tamil": "ta", "Telugu": "te", "Malayalam": "ml", "Kannada": "kn",
              "Marathi": "mr", "Bengali": "bn", "Gujarati": "gu", "Punjabi": "pa", "Odia": "or",
              "Bhojpuri": "bho", "Assamese": "as", "Tulu": "tcy", "Urdu": "ur", "English": "en"}
CODE_LANGS = {v: k for k, v in LANG_CODES.items()}
IMDB_MOVIE_TYPES = ["movie", "tvMovie"]
IMDB_TV_TYPES = ["tvSeries", "tvMiniSeries"]
IMDB_FIELDS = (
    "id titleText{text} titleType{id} releaseDate{day month year} genres{genres{text}} "
    "spokenLanguages{spokenLanguages{text}} countriesOfOrigin{countries{id}} "
    'directors:credits(first:3,filter:{categories:["director"]}){edges{node{name{nameText{text}}}}} '
    'cast:credits(first:3,filter:{categories:["actor","actress"]}){edges{node{name{nameText{text}}}}} '
    'producers:credits(first:3,filter:{categories:["producer"]}){edges{node{name{nameText{text}}}}} '
    'prod:companyCredits(first:3,filter:{categories:["production"]}){edges{node{company{companyText{text}}}}} '
    'dist:companyCredits(first:12,filter:{categories:["distribution"]}){edges{node{company{companyText{text}}}}} '
    "runtime{seconds} certificate{rating}")
IMDB_SEARCH_QUERY = (
    "query($types:[String!]!,$start:Date,$end:Date,$after:String,$first:Int!){"
    "advancedTitleSearch(first:$first,after:$after,constraints:{"
    "titleTypeConstraint:{anyTitleTypeIds:$types},"
    "releaseDateConstraint:{releaseDateRange:{start:$start,end:$end}},"
    'originCountryConstraint:{allCountries:["IN"]}},'
    "sort:{sortBy:RELEASE_DATE,sortOrder:ASC}){total pageInfo{hasNextPage endCursor} "
    "edges{node{title{" + IMDB_FIELDS + "}}}}}")


# --------------------------------------------------------------------------------------------
# small helpers
# --------------------------------------------------------------------------------------------
def norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def resolve_window(mode: str, today: date, custom_start: str = "", custom_end: str = "") -> tuple[date | None, date | None, bool]:
    """Return (from, to, keep_tba) for a run mode."""
    first = today.replace(day=1)
    if mode == "1":
        return date(2024, 1, 1), today + relativedelta(years=1), True
    if mode == "2":
        return first, today + relativedelta(months=3), False
    if mode == "3":
        return first, today + relativedelta(months=6), False
    if mode == "4":
        return None, None, True
    start = _parse_iso(custom_start) or date(2024, 1, 1)
    end = _parse_iso(custom_end) or today + relativedelta(years=1)
    if end < start:
        start, end = end, start
    return start, end, False


def _parse_iso(value: str) -> date | None:
    try:
        return date.fromisoformat((value or "").strip())
    except ValueError:
        return None


def status_for(sort_date: str | None, today: date, future: str = "Confirmed") -> str:
    if not sort_date:
        return "Date TBA"
    return "Released" if sort_date <= today.isoformat() else future


@dataclass
class Title:
    title: str
    category: str = "Movie"
    release_date: str = "TBA (announced)"
    sort_date: str | None = None
    status: str = "Date TBA"
    genre: str = ""
    language: str = "Unspecified"
    produced_by: str = ""
    studio: str = ""
    network: str = ""
    director: str = ""
    cast: str = ""
    trailer_url: str = ""
    youtube_channels: list[str] = field(default_factory=list)
    type_platform: str = "Film – Theatrical"
    certificate_type: str = ""
    certificate_no: str = ""
    duration: str = ""
    cbfc_office: str = ""
    origin: str = "Indian"
    window: str = ""
    imdb_id: str = ""
    source: str = ""
    from_list: bool = False

    def row(self) -> dict[str, str]:
        trailer = self.trailer_url or YT_RESULTS + quote_plus(f"{self.title} official trailer")
        chans = "\n".join(f"{c}|{self.title}" for c in self.youtube_channels[:3])
        net = self.network
        if not net:
            m = re.search(r"(?:OTT|TV)\s*\(([^)]+)\)", self.type_platform or "")
            net = m.group(1) if m else ""
        values = {
            "Title Name": self.title, "IMDb ID": self.imdb_id, "Category": self.category,
            "Release Date": self.release_date, "Release Status": self.status, "Genre": self.genre,
            "Language": self.language, "Produce by": self.produced_by,
            "Studio": self.studio or self.produced_by, "Network": net, "Director": self.director,
            "Cast": self.cast, "Teaser / Trailer on YouTube": trailer, "Top 3 YouTube Channels": chans,
            "Type / Platform": self.type_platform, "Certificate Type": self.certificate_type,
            "Certificate No.": self.certificate_no, "Duration": self.duration,
            "CBFC Office": self.cbfc_office, "Origin": self.origin,
            "OTT Window (Pay 1 / Library)": self.window, "Source": self.source,
        }
        return {k: (v if v else "—") for k, v in values.items()}


MERGE_FIELDS = ("imdb_id", "genre", "produced_by", "studio", "network", "director", "cast",
                "trailer_url", "window", "certificate_type", "certificate_no", "duration", "cbfc_office")


def merge_titles(groups: list[list[Title]]) -> list[Title]:
    """Dedupe across sources; earlier groups win, later ones fill empty fields."""
    kept: list[Title] = []
    by_id: dict[str, Title] = {}
    by_key: dict[tuple[str, str], list[Title]] = {}
    for group in groups:
        for t in group:
            lang = re.split(r"[ (,/]", t.language.strip())[0].lower()
            hit = by_id.get(t.imdb_id) if t.imdb_id else None
            if hit is not None and norm(hit.title) != norm(t.title) and t.category == "TV Show":
                hit = None                        # different seasons of one series stay separate
            if hit is None:
                key = (norm(t.title), (t.sort_date or "")[:4])
                for cand in by_key.get(key, []):
                    clang = re.split(r"[ (,/]", cand.language.strip())[0].lower()
                    if clang == lang or "unspecified" in (clang, lang) or lang in cand.language.lower():
                        hit = cand
                        break
            if hit is not None:
                for fld in MERGE_FIELDS:
                    if not getattr(hit, fld) and getattr(t, fld):
                        setattr(hit, fld, getattr(t, fld))
                if not hit.youtube_channels and t.youtube_channels:
                    hit.youtube_channels = t.youtube_channels
                if t.from_list:
                    hit.from_list = True
                    if t.window:
                        hit.window = t.window
                if t.source and t.source not in hit.source:
                    hit.source = f"{hit.source}; {t.source}" if hit.source else t.source
                if t.imdb_id:
                    by_id.setdefault(t.imdb_id, hit)
                continue
            kept.append(t)
            if t.imdb_id:
                by_id[t.imdb_id] = t
            by_key.setdefault((norm(t.title), (t.sort_date or "")[:4]), []).append(t)
    return kept


def filter_window(titles: list[Title], start: date | None, end: date | None, keep_tba: bool,
                  only_list: bool) -> list[Title]:
    if only_list:
        return [t for t in titles if t.from_list]
    lo = start.isoformat() if start else "0000-01-01"
    hi = end.isoformat() if end else "9999-12-31"
    return [t for t in titles if t.from_list or (lo <= t.sort_date <= hi if t.sort_date else keep_tba)]


# --------------------------------------------------------------------------------------------
# Wikipedia tables
# --------------------------------------------------------------------------------------------
def _cell_text(cell) -> str:
    for tag in cell.find_all(["style", "sup"]):
        tag.decompose()
    for br in cell.find_all("br"):
        br.replace_with(", ")
    for li in cell.find_all("li"):
        li.append(", ")
    text = re.sub(r"\s+", " ", cell.get_text(" "))
    text = re.sub(r"(\s*,\s*)+", ", ", text)
    return text.strip(" ,")


def table_grid(table) -> list[list[str]]:
    """Expand rowspan/colspan so every row has a value in every column."""
    grid: list[list[str | None]] = []
    for ri, tr in enumerate(table.find_all("tr")):
        while len(grid) <= ri:
            grid.append([])
        ci = 0
        for cell in tr.find_all(["td", "th"], recursive=False):
            while ci < len(grid[ri]) and grid[ri][ci] is not None:
                ci += 1
            rs = int(cell.get("rowspan", 1) or 1) if str(cell.get("rowspan", "1")).isdigit() else 1
            cs = int(cell.get("colspan", 1) or 1) if str(cell.get("colspan", "1")).isdigit() else 1
            value = _cell_text(cell)
            for a in range(rs):
                while len(grid) <= ri + a:
                    grid.append([])
                row = grid[ri + a]
                for b in range(cs):
                    while len(row) <= ci + b:
                        row.append(None)
                    row[ci + b] = value
            ci += cs
    return [[c or "" for c in row] for row in grid]


def parse_film_year_page(html: str, language: str, year: int, today: date) -> list[Title]:
    soup = BeautifulSoup(html, "html.parser")
    out: list[Title] = []
    seen: set[str] = set()
    for table in soup.select("table.wikitable"):
        grid = table_grid(table)
        hi = next((i for i, r in enumerate(grid) if any(re.fullmatch(r"(title|film|name)", c.strip(), re.I) for c in r)), -1)
        if hi < 0:
            continue
        head = [c.lower() for c in grid[hi]]
        if any(re.search(r"rank|gross|box office|collection|award|category", h) for h in head):
            continue
        col = lambda pat: next((i for i, h in enumerate(head) if re.search(pat, h)), -1)  # noqa: E731
        c_t = next(i for i, h in enumerate(head) if re.fullmatch(r"(title|film|name)", h.strip()))
        c_d, c_c = col(r"director"), col(r"cast|starring")
        c_p, c_g, c_r = col(r"production|studio|producer|banner"), col(r"genre"), col(r"release date|^date")
        for row in grid[hi + 1:]:
            if len(row) <= c_t or not row[c_t] or re.fullmatch(r"(title|film)", row[c_t], re.I):
                continue
            title = row[c_t]
            if norm(title) in seen:
                continue
            seen.add(norm(title))
            mon = day = None
            cells = row[:c_t] + ([row[c_r]] if 0 <= c_r < len(row) and c_r > c_t else [])
            for raw in cells:
                v = re.sub(r"[^a-z0-9]", "", raw.lower())
                if v[:3] in MONTHS and v.isalpha():
                    mon = MONTHS[v[:3]]
                elif v.isdigit() and len(v) <= 2:
                    day = int(v)
                else:
                    m = re.match(r"^(\d{1,2})([a-z]{3,9})|^([a-z]{3,9})(\d{1,2})", v)
                    if m and (m.group(2) or m.group(3))[:3] in MONTHS:
                        mon = MONTHS[(m.group(2) or m.group(3))[:3]]
                        day = int(m.group(1) or m.group(4))
            get = lambda i: row[i] if 0 <= i < len(row) else ""  # noqa: E731
            prod = get(c_p)
            t = Title(title=title, language=language, director=get(c_d)[:80],
                      cast=", ".join(get(c_c).split(", ")[:3]), produced_by=prod, genre=get(c_g),
                      source=f"Wikipedia – List of {language} films of {year}")
            net = NETWORK_RE.search(prod or "")
            if net:
                t.network = net.group(0).strip()
                t.studio = re.sub(r",?\s*(Netflix|Amazon Prime Video|Prime Video|JioHotstar|ZEE5|SonyLIV|Sun NXT|ETV Win)\b", "", prod).strip(" ,/")
            try:
                if mon and day:
                    d = date(year, mon, day)
                    t.release_date, t.sort_date = d.strftime("%d %b %Y"), d.isoformat()
                elif mon:
                    d = date(year, mon, 1)
                    t.release_date, t.sort_date = d.strftime("%b %Y"), d.isoformat()
            except ValueError:
                pass
            t.status = status_for(t.sort_date, today)
            out.append(t)
    return out


def parse_tv_list_page(html: str, platform: str, today: date) -> list[Title]:
    soup = BeautifulSoup(html, "html.parser")
    out: list[Title] = []
    for table in soup.select("table.wikitable"):
        grid = table_grid(table)
        if not grid:
            continue
        hi = next((i for i, r in enumerate(grid) if any(re.fullmatch(r"(title|series|show)", c.strip(), re.I) for c in r)), -1)
        if hi < 0:
            continue
        head = [c.lower() for c in grid[hi]]
        c_t = next(i for i, h in enumerate(head) if re.fullmatch(r"(title|series|show)", h.strip()))
        c_d = next((i for i, h in enumerate(head) if re.search(r"premiere|release|date", h)), -1)
        c_l = next((i for i, h in enumerate(head) if "language" in h), -1)
        c_g = next((i for i, h in enumerate(head) if "genre" in h), -1)
        last_lang = ""
        for row in grid[hi + 1:]:
            if len(row) <= c_t or not row[c_t]:
                continue
            raw_date = row[c_d] if 0 <= c_d < len(row) else ""
            d = None
            if raw_date:
                try:
                    d = dateparser.parse(re.sub(r"\(.*?\)|\[.*?\]", "", raw_date), default=datetime(1900, 1, 1)).date()
                except (ValueError, OverflowError):
                    d = None
            if d is None or d.year < 2024:
                continue
            lang = (row[c_l] if 0 <= c_l < len(row) else "") or last_lang
            last_lang = lang
            t = Title(title=row[c_t], category="TV Show", language=lang or "Hindi",
                      genre=row[c_g] if 0 <= c_g < len(row) else "", network=platform,
                      type_platform=f"Web series – OTT ({platform})",
                      release_date=d.strftime("%d %b %Y"), sort_date=d.isoformat(),
                      source=f"Wikipedia – {platform} originals list")
            t.status = status_for(t.sort_date, today)
            out.append(t)
    return out


# --------------------------------------------------------------------------------------------
# YouTube
# --------------------------------------------------------------------------------------------
def _walk_videos(obj, out: list) -> None:
    if isinstance(obj, dict):
        if "videoRenderer" in obj:
            out.append(obj["videoRenderer"])
        for v in obj.values():
            _walk_videos(v, out)
    elif isinstance(obj, list):
        for v in obj:
            _walk_videos(v, out)


def parse_youtube_results(html: str) -> list[dict[str, Any]] | None:
    """Videos from a YouTube results page; None when the page has no ytInitialData (blocked)."""
    m = re.search(r"var ytInitialData = (\{.*?\});</script>", html, re.S)
    if not m:
        return None
    vids: list = []
    _walk_videos(json.loads(m.group(1)), vids)
    out = []
    for v in vids:
        run = ((v.get("ownerText") or {}).get("runs") or [{}])[0] or ((v.get("longBylineText") or {}).get("runs") or [{}])[0]
        ep = ((run.get("navigationEndpoint") or {}).get("browseEndpoint") or {})
        if not ep or not v.get("videoId"):
            continue
        out.append({
            "id": v["videoId"],
            "t": "".join(r.get("text", "") for r in (v.get("title") or {}).get("runs", [])),
            "ch": "https://www.youtube.com" + (ep.get("canonicalBaseUrl") or "/channel/" + ep.get("browseId", "")),
            "ver": "VERIFIED" in json.dumps(v.get("ownerBadges") or []),
        })
    return out


YT_ROLES = (("trailer", re.compile(r"trailer|teaser|glimpse", re.I), True),
            ("full movie", re.compile(r"full\s*movie|full\s*hd|movie|episode", re.I), False),
            ("songs", re.compile(r"song|lyric|jukebox|video", re.I), False))


def pick_top3(title: str, role_results: list[list[dict[str, Any]]]) -> tuple[str, list[str]]:
    """One channel per role (trailer, full movie, songs), verified first, then fill from the rest."""
    base = re.sub(r"\s*\(.*?\)\s*", " ", title)
    base = re.sub(r"\s*[-–]?\s*season\s*\d+.*$", "", base, flags=re.I).strip()
    key = norm(base)
    picked: list[str] = []
    pool: list[dict[str, Any]] = []
    trailer = ""
    for (_, pattern, is_trailer), results in zip(YT_ROLES, role_results):
        matches = [x for x in results if key and key in norm(x["t"])]
        pool.extend(matches)
        cands = [x for x in matches if pattern.search(x["t"]) and x["ch"] not in picked]
        choice = next((x for x in cands if x["ver"]), cands[0] if cands else None)
        if choice:
            picked.append(choice["ch"])
            if is_trailer:
                trailer = "https://www.youtube.com/watch?v=" + choice["id"]
    for x in [x for x in pool if x["ver"]] + pool:
        if len(picked) < 3 and x["ch"] not in picked:
            picked.append(x["ch"])
    return trailer, picked[:3]


# --------------------------------------------------------------------------------------------
# CBFC website
# --------------------------------------------------------------------------------------------
CBFC_LABELS = {"film name": "film", "movie language": "language", "movie category": "category",
               "certified by regional office": "office", "certificate no": "certno",
               "certificate date": "certdate", "certified length": "length", "name of producer": "producer"}


def parse_cbfc_results(html: str) -> list[dict[str, str]]:
    lines = [l.strip() for l in BeautifulSoup(html, "html.parser").get_text("\n").split("\n") if l.strip()]
    recs: list[dict[str, str]] = []
    cur: dict[str, str] | None = None
    for i, line in enumerate(lines):
        key = CBFC_LABELS.get(re.sub(r"[:.]", "", line.lower()).strip())
        if not key:
            continue
        if key == "film":
            cur = {}
            recs.append(cur)
        if cur is not None:
            nxt = lines[i + 1] if i + 1 < len(lines) else ""
            cur[key] = "" if CBFC_LABELS.get(re.sub(r"[:.]", "", nxt.lower()).strip()) else nxt
    return [r for r in recs if r.get("film") and r["film"] != "-"]


# --------------------------------------------------------------------------------------------
# uploaded lists
# --------------------------------------------------------------------------------------------
def read_title_list(file_content: bytes | None, filename: str, bulk_text: str) -> list[dict[str, Any]]:
    rows: list[list[Any]] = []
    if file_content and filename:
        if filename.lower().endswith(".csv"):
            rows = list(csv.reader(StringIO(file_content.decode("utf-8-sig", errors="replace"))))
        else:
            from openpyxl import load_workbook
            ws = load_workbook(BytesIO(file_content), data_only=True, read_only=True).active
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
    out: list[dict[str, Any]] = []
    if rows:
        head = [str(h or "").strip().lower() for h in rows[0]]
        col = lambda *keys: next((i for i, h in enumerate(head) if any(k in h for k in keys)), None)  # noqa: E731
        c = dict(imdb=col("imdb"), title=col("title"), lang=col("language", "lang"), date=col("date", "year"),
                 ctype=col("content type", "type", "category"), window=col("pay1", "pay 1", "library", "window"))
        if c["title"] is None:
            c["title"] = 0
        for r in rows[1:]:
            g = lambda k: r[c[k]] if c[k] is not None and c[k] < len(r) else None  # noqa: E731
            if g("title"):
                out.append({k: g(k) for k in c})
    for line in (bulk_text or "").splitlines():
        parts = [p.strip() for p in line.split("|")]
        if parts and parts[0]:
            out.append({"title": parts[0], "lang": parts[1] if len(parts) > 1 else "",
                        "date": parts[2] if len(parts) > 2 else "", "imdb": parts[3] if len(parts) > 3 else "",
                        "ctype": "", "window": ""})
    return out


def to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:       # Excel serial
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = clean(value)
    if re.fullmatch(r"\d{4}", text):
        return date(int(text), 12, 31)
    if text:
        try:
            return dateparser.parse(text, dayfirst=True).date()
        except (ValueError, OverflowError):
            return None
    return None


# --------------------------------------------------------------------------------------------
# the service
# --------------------------------------------------------------------------------------------
class IndianReleaseTrackerService:
    def __init__(self, http_client, youtube_api_key: str = "", tmdb_api_key: str = "",
                 tmdb_read_access_token: str = "", omdb_api_key: str = "") -> None:
        self.http = http_client
        self.youtube_api_key = youtube_api_key
        self.tmdb_api_key = tmdb_api_key
        self.tmdb_token = tmdb_read_access_token
        self.omdb_api_key = omdb_api_key
        self.notes: list[str] = []

    # -- HTTP --------------------------------------------------------------------------------
    def _get_json(self, url: str, params: dict[str, Any] | None = None, headers: dict | None = None) -> dict:
        r = self.http.session.get(url, params=params, headers=headers, timeout=self.http.timeout_seconds)
        r.raise_for_status()
        return r.json()

    def _wiki_page_html(self, page: str) -> str | None:
        try:
            data = self._get_json(WIKI_API, {"action": "parse", "prop": "text", "redirects": 1,
                                            "format": "json", "page": page})
        except Exception:
            return None
        return (data.get("parse") or {}).get("text", {}).get("*")

    # -- run ---------------------------------------------------------------------------------
    def run(self, *, mode: str, custom_start: str = "", custom_end: str = "", languages: list[str] | None = None,
            include_tv: bool = True, youtube_scope: str = "recent", check_cbfc: bool = True,
            file_content: bytes | None = None, filename: str = "", bulk_text: str = "",
            api_key_override: str = "", notes: str = "", progress: ProgressCallback | None = None,
            today: date | None = None) -> dict[str, Any]:
        progress = progress or (lambda p, m: None)
        today = today or date.today()
        self.notes = []
        languages = [l for l in (languages or INDIAN_LANGUAGES) if l in INDIAN_LANGUAGES] or INDIAN_LANGUAGES
        start, end, keep_tba = resolve_window(mode, today, custom_start, custom_end)
        only_list = mode == "4"
        listed = read_title_list(file_content, filename, bulk_text)
        if only_list and not listed:
            raise ValueError("Option 4 needs a title list: upload a CSV/Excel file or paste titles.")

        groups: list[list[Title]] = []
        coverage: list[dict[str, str]] = []
        if not only_list:
            years = list(range(start.year, end.year + 1))
            groups.append(self._imdb_advanced_search(start, end, include_tv, languages, today, progress))
            groups.append(self._wikipedia_films(languages, years, today, progress))
            if include_tv:
                groups.append(self._wikipedia_tv(today, progress))
        if listed:
            list_titles, coverage = self._resolve_list(listed, today, progress)
            groups.append(list_titles)

        progress(55, "Merging sources")
        titles = filter_window(merge_titles(groups), start, end, keep_tba, only_list)
        titles.sort(key=lambda t: (t.sort_date or "9999-12-31", t.title.lower()))

        self._tmdb_enrich(titles, today, progress)
        self._omdb_enrich(titles, progress)

        self._youtube(titles, youtube_scope, today, api_key_override or self.youtube_api_key, progress)
        if check_cbfc:
            self._cbfc(titles, today, progress)

        progress(97, "Building result tables")
        rows = [t.row() for t in titles]
        window_text = ("your title list only" if only_list else
                       f"{start:%d %b %Y} → {end:%d %b %Y}" + (" + Date TBA titles" if keep_tba else ""))
        movies = sum(t.category == "Movie" for t in titles)
        with_yt = sum(bool(t.youtube_channels) for t in titles)
        with_net = sum(bool(r["Network"] != "—") for r in rows)
        by_status = ", ".join(f"{s}: {sum(t.status == s for t in titles)}" for s in STATUSES if any(t.status == s for t in titles))
        summary = (f"Mode {mode} – {window_text}. {len(titles)} titles ({movies} movies, {len(titles) - movies} TV shows); "
                   f"{sum(bool(t.imdb_id) for t in titles)} with an IMDb ID, {with_yt} with Top 3 YouTube channels, "
                   f"{with_net} with a Network. {by_status}.")
        if notes:
            summary += f" Custom notes: {clean(notes)}."
        if self.notes:
            summary += " Notes: " + " ".join(self.notes)
        movie_rows = [r for r in rows if r["Category"] == "Movie"]
        tv_rows = [r for r in rows if r["Category"] != "Movie"]
        # Two tables so each gets its own BDR ingest template (Movies / TV Shows) - the ingest
        # export maps a section's tracker_type to one title category.
        sections = [{"key": "movies", "title": "Indian Regional Movies", "columns": TITLE_COLUMNS,
                     "rows": movie_rows, "row_count": len(movie_rows), "supports_google": True,
                     "tracker_type": TRACKER_TYPE + "_movies"}]
        if tv_rows or not only_list:
            sections.append({"key": "tv_shows", "title": "Indian Regional TV Shows", "columns": TITLE_COLUMNS,
                             "rows": tv_rows, "row_count": len(tv_rows), "supports_google": True,
                             "tracker_type": TRACKER_TYPE + "_tv"})
        if coverage:
            sections.append({"key": "source_coverage", "title": "Source Coverage (uploaded list)",
                             "columns": COVERAGE_COLUMNS, "rows": coverage, "row_count": len(coverage),
                             "supports_google": True})
        return {"tracker_type": TRACKER_TYPE, "title": "Indian Regional Release Tracker",
                "created_at": utc_now_iso(), "source_url": "https://en.wikipedia.org/wiki/List_of_Indian_films",
                "summary": summary, "sections": sections}

    # -- self-test (GET /indian-release-tracker/diagnostics) ---------------------------------
    def diagnose(self) -> list[dict[str, str]]:
        """Probe every source once FROM THIS SERVER, so a deployment shows what it can reach."""
        today = date.today()
        checks: list[dict[str, str]] = []

        def probe(name: str, fn) -> None:
            t0 = time.time()
            try:
                ok, detail = fn()
            except Exception as exc:  # noqa: BLE001 - report, never raise
                ok, detail = False, f"{exc.__class__.__name__}: {str(exc)[:160]}"
            checks.append({"source": name, "status": "OK" if ok else ("SKIPPED" if ok is None else "BLOCKED / FAILED"),
                           "seconds": f"{time.time() - t0:.1f}", "detail": detail})

        def imdb():
            r = self.http.session.post(IMDB_GRAPHQL, timeout=self.http.timeout_seconds,
                                       headers={"content-type": "application/json"},
                                       json={"query": IMDB_SEARCH_QUERY, "variables": {
                                           "types": IMDB_MOVIE_TYPES + IMDB_TV_TYPES, "start": today.isoformat(),
                                           "end": (today + timedelta(days=30)).isoformat(), "after": None, "first": 3}})
            data = (r.json().get("data") or {}).get("advancedTitleSearch") if r.status_code == 200 else None
            if not data:
                return False, f"HTTP {r.status_code}, no advancedTitleSearch data"
            first = [e["node"]["title"]["titleText"]["text"] for e in data.get("edges") or []]
            return True, f"{data.get('total')} India titles in the next 30 days, e.g. {', '.join(first)}"

        def tmdb():
            if not (self.tmdb_api_key or self.tmdb_token):
                return None, "TMDB_API_KEY / TMDB_READ_ACCESS_TOKEN not set"
            found = self._tmdb("/find/tt0321968", {"external_source": "imdb_id"})
            hit = (found.get("movie_results") or [{}])[0]
            return bool(hit), f"tt0321968 -> {hit.get('title', 'no match')} ({hit.get('release_date', '')})"

        def omdb():
            if not self.omdb_api_key:
                return None, "OMDB_API_KEY not set"
            d = self._get_json(OMDB_API, {"i": "tt0321968", "apikey": self.omdb_api_key})
            return d.get("Response") == "True", f"{d.get('Title') or d.get('Error')} – {d.get('Director', '')}"

        def youtube_page():
            vids = parse_youtube_results(self.http.get_text(YT_RESULTS + quote_plus("Deep Jele Jai Bengali trailer")))
            if vids is None:
                return False, "results page served without ytInitialData (consent/bot page)"
            return bool(vids), f"{len(vids)} videos; first channel {vids[0]['ch'] if vids else '-'}"

        def youtube_api():
            if not self.youtube_api_key:
                return None, "YOUTUBE_API_KEY not set (only used if the results page is blocked)"
            res = self._youtube_api_search("Deep Jele Jai Bengali trailer", self.youtube_api_key)
            return bool(res), f"{len(res)} videos; first channel {res[0]['ch'] if res else '-'}"

        def wikipedia():
            html = self._wiki_page_html(f"List of Hindi films of {today.year}")
            rows = parse_film_year_page(html, "Hindi", today.year, today) if html else []
            return bool(rows), f"List of Hindi films of {today.year}: {len(rows)} films parsed"

        def wikidata():
            d = self._get_json(WIKIDATA_SPARQL, {"format": "json", "query":
                               'SELECT ?item WHERE { ?item wdt:P345 "tt0321968" }'},
                               headers={"Accept": "application/sparql-results+json"})
            b = d["results"]["bindings"]
            return bool(b), f"tt0321968 -> {b[0]['item']['value'].rsplit('/', 1)[-1] if b else 'no item'}"

        def cbfc():
            form = self._cbfc_form()
            return (form is not None), (f"Search Film form found: {form[1].upper()} {form[0]} (field '{form[3]}')"
                                        if form else "no Search Film form served to this server")

        for name, fn in (("IMDb Advanced Search (GraphQL)", imdb), ("TMDB API", tmdb), ("OMDb API", omdb),
                         ("YouTube results page", youtube_page), ("YouTube Data API", youtube_api),
                         ("Wikipedia API", wikipedia), ("Wikidata SPARQL", wikidata), ("CBFC website", cbfc)):
            probe(name, fn)
        return checks

    # -- sources -----------------------------------------------------------------------------
    def _wikipedia_films(self, languages, years, today, progress) -> list[Title]:
        out: list[Title] = []
        jobs = [(l, y) for l in languages for y in years]
        missing = []
        for i, (lang, year) in enumerate(jobs):
            progress(10 + int(30 * i / max(len(jobs), 1)), f"Wikipedia: {lang} films of {year}")
            html = self._wiki_page_html(f"List of {lang} films of {year}") or \
                self._wiki_page_html(f"List of Indian {lang} films of {year}")
            if not html:
                missing.append(f"{lang} {year}")
                continue
            out.extend(parse_film_year_page(html, lang, year, today))
        if missing:
            self.notes.append(f"No Wikipedia year page yet for: {', '.join(missing[:12])}{'…' if len(missing) > 12 else ''}.")
        return out

    def _wikipedia_tv(self, today, progress) -> list[Title]:
        out: list[Title] = []
        for page, platform in TV_LISTS:
            progress(42, f"Wikipedia: {platform} originals")
            html = self._wiki_page_html(page)
            if html:
                out.extend(parse_tv_list_page(html, platform, today))
        return out

    def _imdb_advanced_search(self, start, end, include_tv, languages, today, progress) -> list[Title]:
        """IMDb Advanced Title Search: India-origin titles released in the window (movies + TV)."""
        kinds = [("Movie", IMDB_MOVIE_TYPES)] + ([("TV Show", IMDB_TV_TYPES)] if include_tv else [])
        wanted = {l.lower() for l in languages}
        out: list[Title] = []
        for category, types in kinds:
            after, pages, total = None, 0, 0
            while pages < IMDB_SEARCH_MAX_PAGES:
                progress(5 + pages, f"IMDb advanced search: {category}s {start:%d %b %Y} → {end:%d %b %Y}")
                try:
                    r = self.http.session.post(
                        IMDB_GRAPHQL, timeout=self.http.timeout_seconds,
                        headers={"content-type": "application/json"},
                        json={"query": IMDB_SEARCH_QUERY, "variables": {
                            "types": types, "start": start.isoformat(), "end": end.isoformat(),
                            "after": after, "first": IMDB_SEARCH_PAGE}})
                    data = (r.json().get("data") or {}).get("advancedTitleSearch")
                except Exception as exc:
                    self.notes.append(f"IMDb advanced search not reachable from the server ({exc.__class__.__name__}); "
                                      "Wikipedia and TMDB cover these titles instead.")
                    return out
                if not data:
                    self.notes.append("IMDb advanced search returned no data (IMDb may be blocking the server).")
                    return out
                total = data.get("total") or 0
                for edge in data.get("edges") or []:
                    t = self._imdb_title(edge["node"]["title"], category, wanted, today)
                    if t:
                        out.append(t)
                pages += 1
                info = data.get("pageInfo") or {}
                if not info.get("hasNextPage"):
                    break
                after = info.get("endCursor")
            if pages >= IMDB_SEARCH_MAX_PAGES:
                self.notes.append(f"IMDb advanced search for {category}s stopped at {pages * IMDB_SEARCH_PAGE} of {total} titles.")
        return out

    @staticmethod
    def _imdb_title(t: dict, category: str, wanted: set[str], today: date) -> Title | None:
        names = lambda key, path: [e["node"][path[0]][path[1]]["text"]  # noqa: E731
                                   for e in ((t.get(key) or {}).get("edges") or [])]
        langs = [x["text"] for x in ((t.get("spokenLanguages") or {}).get("spokenLanguages") or [])]
        indian = [l for l in langs if l.lower() in INDIAN_LOWER]
        if indian and wanted and not any(l.lower() in wanted for l in indian):
            return None                                   # a language the user unticked
        rd = t.get("releaseDate") or {}
        d = None
        if rd.get("year"):
            try:
                d = date(rd["year"], rd.get("month") or 1, rd.get("day") or 1)
            except ValueError:
                d = None
        release = (d.strftime("%d %b %Y") if rd.get("day") else d.strftime("%b %Y")) if d else "TBA (announced)"
        dist = [n for n in names("dist", ("company", "companyText")) if NETWORK_RE.search(n)]
        prod = ", ".join(names("prod", ("company", "companyText")))
        type_id = (t.get("titleType") or {}).get("id", "")
        runtime = ((t.get("runtime") or {}).get("seconds") or 0) // 60
        net = ", ".join(dict.fromkeys(dist[:3]))
        if category == "Movie":
            platform = "TV Movie" if type_id == "tvMovie" else "Film – Theatrical"
        else:
            platform = f"{'Mini-series' if type_id == 'tvMiniSeries' else 'TV series'} – " + (f"OTT ({dist[0]})" if dist else "TV / OTT")
        countries = [x["id"] for x in ((t.get("countriesOfOrigin") or {}).get("countries") or [])]
        return Title(
            title=(t.get("titleText") or {}).get("text", ""), imdb_id=t["id"], category=category,
            genre=" / ".join(g["text"] for g in ((t.get("genres") or {}).get("genres") or [])),
            language=(langs[0] + (f" (+ {', '.join(langs[1:])})" if langs[1:] else "")) if langs else "Unspecified",
            director=", ".join(names("directors", ("name", "nameText"))),
            cast=", ".join(names("cast", ("name", "nameText"))),
            produced_by=", ".join(names("producers", ("name", "nameText"))) or prod, studio=prod, network=net,
            release_date=release, sort_date=d.isoformat() if d else None,
            status=status_for(d.isoformat() if d else None, today, "Tentative"),
            duration=f"{runtime} min" if runtime else "",
            certificate_type=((t.get("certificate") or {}).get("rating") or ""),
            type_platform=platform,
            origin="Indian" if ("IN" in countries and (indian or not langs)) else "International",
            source="IMDb Advanced Search (country of origin: India)")

    # -- TMDB / OMDb (the hub's API keys) ----------------------------------------------------
    def _tmdb(self, path: str, params: dict | None = None) -> dict:
        params = dict(params or {})
        headers = {"Accept": "application/json"}
        if self.tmdb_token:
            headers["Authorization"] = f"Bearer {self.tmdb_token}"
        elif self.tmdb_api_key:
            params["api_key"] = self.tmdb_api_key
        return self._get_json(TMDB_API + path, params, headers)

    def _tmdb_enrich(self, titles: list[Title], today: date, progress) -> None:
        if not (self.tmdb_api_key or self.tmdb_token):
            self.notes.append("TMDB key not configured – skipped TMDB enrichment.")
            return
        needs = [t for t in titles if not (t.imdb_id and t.genre and t.director and t.studio and t.network)]
        needs.sort(key=lambda t: (not t.from_list, t.sort_date is None, -(int(t.sort_date.replace("-", "")) if t.sort_date else 0)))
        targets = needs[:TMDB_MAX_TITLES]
        if len(needs) > TMDB_MAX_TITLES:
            self.notes.append(f"TMDB enrichment capped at {TMDB_MAX_TITLES} of {len(needs)} titles (listed, upcoming and recent first).")
        if not targets:
            return
        state = {"done": 0, "hit": 0}

        def one(t: Title) -> None:
            try:
                kind, tid = self._tmdb_locate(t)
                if tid:
                    self._apply_tmdb(t, kind, tid)
                    state["hit"] += 1
            except Exception:
                pass
            state["done"] += 1
            if state["done"] % 25 == 0:
                progress(56 + int(4 * state["done"] / len(targets)), f"TMDB: {state['done']}/{len(targets)} titles")

        progress(56, f"TMDB: enriching {len(targets)} titles")
        with ThreadPoolExecutor(max_workers=TMDB_WORKERS) as pool:
            list(pool.map(one, targets))
        self.notes.append(f"TMDB matched {state['hit']} of {len(targets)} titles that had gaps.")

    def _tmdb_locate(self, t: Title) -> tuple[str, int | None]:
        want = "tv" if t.category != "Movie" else "movie"
        if t.imdb_id:
            found = self._tmdb(f"/find/{t.imdb_id}", {"external_source": "imdb_id"})
            for kind, key in (("movie", "movie_results"), ("tv", "tv_results")):
                if found.get(key):
                    return kind, found[key][0]["id"]
            return want, None
        lang = LANG_CODES.get(t.language.split(" (")[0].split()[0] if t.language else "", "")
        params = {"query": re.sub(r"\s*\(.*?\)\s*", " ", t.title).strip()}
        if t.sort_date:
            params["year" if want == "movie" else "first_air_date_year"] = t.sort_date[:4]
        results = self._tmdb(f"/search/{want}", params).get("results") or []
        key = norm(params["query"])
        for r in results[:8]:
            names = [r.get("title") or r.get("name") or "", r.get("original_title") or r.get("original_name") or ""]
            if any(norm(n) == key for n in names) and (not lang or r.get("original_language") == lang):
                return want, r["id"]
        return want, None

    def _apply_tmdb(self, t: Title, kind: str, tid: int) -> None:
        d = self._tmdb(f"/{kind}/{tid}", {"append_to_response": "credits,external_ids,watch/providers"})
        ext = d.get("external_ids") or {}
        if not t.imdb_id and (ext.get("imdb_id") or d.get("imdb_id") or "").startswith("tt"):
            t.imdb_id = ext.get("imdb_id") or d.get("imdb_id")
        if not t.genre:
            t.genre = " / ".join(g["name"] for g in d.get("genres") or [])
        credits = d.get("credits") or {}
        if not t.director:
            crew = credits.get("crew") or []
            t.director = ", ".join(dict.fromkeys(c["name"] for c in crew if c.get("job") in ("Director", "Series Director")))[:120] \
                or ", ".join(c.get("name", "") for c in (d.get("created_by") or []))
        if not t.cast:
            t.cast = ", ".join(c["name"] for c in (credits.get("cast") or [])[:3])
        if not t.studio:
            t.studio = ", ".join(c["name"] for c in (d.get("production_companies") or [])[:3])
        if not t.network:
            nets = [n["name"] for n in d.get("networks") or []]
            providers = ((d.get("watch/providers") or {}).get("results") or {}).get("IN") or {}
            flat = [p["provider_name"] for p in providers.get("flatrate") or []]
            t.network = ", ".join(dict.fromkeys((nets or flat)[:3]))
        if t.language in ("", "Unspecified") and d.get("original_language") in CODE_LANGS:
            t.language = CODE_LANGS[d["original_language"]]
        if not t.sort_date:
            raw = d.get("release_date") or d.get("first_air_date") or ""
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", raw):
                dd = date.fromisoformat(raw)
                t.sort_date, t.release_date = raw, dd.strftime("%d %b %Y")
        t.source += "; TMDB"

    def _omdb_enrich(self, titles: list[Title], progress) -> None:
        if not self.omdb_api_key:
            return
        targets = [t for t in titles if t.imdb_id and not (t.genre and t.director and t.cast)][:OMDB_MAX_TITLES]
        if not targets:
            return
        progress(59, f"OMDb: filling gaps for {len(targets)} titles")
        filled = 0
        for t in targets:
            try:
                d = self._get_json(OMDB_API, {"i": t.imdb_id, "apikey": self.omdb_api_key})
            except Exception:
                continue
            if d.get("Response") != "True":
                continue
            ok = lambda v: v and v != "N/A"  # noqa: E731
            if not t.genre and ok(d.get("Genre")):
                t.genre = d["Genre"].replace(", ", " / ")
            if not t.director and ok(d.get("Director")):
                t.director = d["Director"]
            if not t.cast and ok(d.get("Actors")):
                t.cast = d["Actors"]
            filled += 1
            t.source += "; OMDb"
        if filled:
            self.notes.append(f"OMDb filled gaps for {filled} titles.")

    def _resolve_list(self, listed, today, progress) -> tuple[list[Title], list[dict[str, str]]]:
        progress(45, f"Resolving {len(listed)} listed titles on Wikidata")
        ids = sorted({clean(r.get("imdb")) for r in listed if re.fullmatch(r"tt\d+", clean(r.get("imdb")))})
        info: dict[str, dict] = {}
        for i in range(0, len(ids), 150):
            values = " ".join(f'"{x}"' for x in ids[i:i + 150])
            query = ("SELECT ?imdb ?item (GROUP_CONCAT(DISTINCT ?cc;separator='/') AS ?co) "
                     "(GROUP_CONCAT(DISTINCT ?ll;separator='/') AS ?lang) (GROUP_CONCAT(DISTINCT ?dl;separator=', ') AS ?dir) "
                     "(GROUP_CONCAT(DISTINCT ?sl;separator=', ') AS ?studio) (GROUP_CONCAT(DISTINCT ?nl;separator=', ') AS ?net) "
                     f"WHERE {{ VALUES ?imdb {{{values}}} ?item wdt:P345 ?imdb. "
                     "OPTIONAL{?item wdt:P495/wdt:P297 ?cc.} OPTIONAL{?item wdt:P364 ?l. ?l rdfs:label ?ll FILTER(lang(?ll)='en')} "
                     "OPTIONAL{?item wdt:P57 ?d. ?d rdfs:label ?dl FILTER(lang(?dl)='en')} "
                     "OPTIONAL{?item wdt:P272 ?s. ?s rdfs:label ?sl FILTER(lang(?sl)='en')} "
                     "OPTIONAL{?item wdt:P449 ?n. ?n rdfs:label ?nl FILTER(lang(?nl)='en')} } GROUP BY ?imdb ?item")
            try:
                data = self._get_json(WIKIDATA_SPARQL, {"format": "json", "query": query},
                                      headers={"Accept": "application/sparql-results+json"})
                for b in data["results"]["bindings"]:
                    info.setdefault(b["imdb"]["value"], {k: v["value"] for k, v in b.items()})
            except Exception as exc:
                self.notes.append(f"Wikidata lookup failed ({exc.__class__.__name__}).")
                break
        titles, coverage = [], []
        for r in listed:
            title = clean(r.get("title"))
            iid = clean(r.get("imdb")) if re.fullmatch(r"tt\d+", clean(r.get("imdb"))) else ""
            d = to_date(r.get("date"))
            lang = clean(r.get("lang")).title() or "Unspecified"
            ctype = clean(r.get("ctype")).lower()
            cat = "Movie" if (not ctype or "movie" in ctype or "film" in ctype) else "TV Show"
            b = info.get(iid, {})
            countries = (b.get("co") or "").split("/")
            orig_langs = [x.replace(" language", "") for x in (b.get("lang") or "").split("/") if x]
            origin = "Indian" if ("IN" in countries or lang.lower().split()[0] in INDIAN_LOWER) else "International"
            note = []
            language = lang
            if orig_langs and lang.lower().split()[0] not in " ".join(orig_langs).lower():
                language = f"{lang} (original: {', '.join(orig_langs)})"
                note.append(f"{lang} looks like a dub; original {', '.join(orig_langs)}")
            qid = (b.get("item") or "").rsplit("/", 1)[-1]
            srcs = ["Wikidata " + qid] if qid else []
            if iid:
                srcs.append("IMDb")
            srcs.append("Uploaded title list")
            if not qid:
                note.append("Not on Wikidata – check Sacnilk / JustWatch India by title")
            sd = d.isoformat() if d else None
            titles.append(Title(
                title=title, imdb_id=iid, category=cat, language=language, director=b.get("dir", ""),
                studio=b.get("studio", ""), network=b.get("net", ""), origin=origin,
                window=clean(r.get("window")), release_date=d.strftime("%d %b %Y") if d else "TBA",
                sort_date=sd, status=status_for(sd, today), from_list=True, source="; ".join(srcs),
                type_platform="Film – OTT" if cat == "Movie" else "TV series – OTT"))
            coverage.append({"Title": title, "IMDb ID": iid or "—", "Language": lang, "Year": str(d.year) if d else "",
                             "Origin": origin, "Window": clean(r.get("window")) or "—", "Wikidata": qid or "—",
                             "Sources found": "; ".join(srcs), "Notes": "; ".join(note)})
        return titles, coverage

    def _youtube(self, titles, scope, today, api_key, progress) -> None:
        if scope == "none":
            return
        if scope == "recent":
            cutoff = (today - timedelta(days=90)).isoformat()
            targets = [t for t in titles if t.from_list or not t.sort_date or t.sort_date >= cutoff]
        else:
            targets = list(titles)
        if len(targets) > YOUTUBE_MAX_TITLES:
            self.notes.append(f"YouTube lookup capped at {YOUTUBE_MAX_TITLES} of {len(targets)} titles (upcoming/list titles first).")
            targets = sorted(targets, key=lambda t: (not t.from_list, -(int(t.sort_date.replace('-', '')) if t.sort_date else 99999999)))[:YOUTUBE_MAX_TITLES]
        if not targets:
            return
        state = {"done": 0, "blocked": False, "api_used": 0}

        def search_page(q: str) -> list[dict[str, Any]] | None:
            try:
                return parse_youtube_results(self.http.get_text(YT_RESULTS + quote_plus(q)))
            except Exception:
                return None

        def one(t: Title) -> None:
            lang = t.language.split(" (")[0].split()[0] if t.language and t.language != "Unspecified" else ""
            base = re.sub(r"\s*\(.*?\)\s*", " ", t.title).strip()
            queries = [f"{base} {lang} {role}".replace("  ", " ") for role, _, _ in YT_ROLES]
            results = [None if state["blocked"] else search_page(q) for q in queries]
            if any(r is None for r in results):
                state["blocked"] = True
                if api_key and state["api_used"] < YOUTUBE_API_MAX_TITLES:
                    state["api_used"] += 1
                    results = [self._youtube_api_search(q, api_key) for q in queries]
                else:
                    return
            trailer, chans = pick_top3(t.title, results)
            if chans:
                t.youtube_channels = chans
            if trailer and not t.trailer_url:
                t.trailer_url = trailer
            state["done"] += 1
            if state["done"] % 10 == 0:
                progress(60 + int(25 * state["done"] / len(targets)), f"YouTube: {state['done']}/{len(targets)} titles")

        progress(60, f"YouTube: top 3 channels for {len(targets)} titles")
        with ThreadPoolExecutor(max_workers=YOUTUBE_WORKERS) as pool:
            list(pool.map(one, targets))
        if state["blocked"]:
            self.notes.append("YouTube's results page was blocked for the server"
                              + (f"; used the YouTube Data API for {state['api_used']} titles (quota-limited)." if state["api_used"]
                                 else "; add YOUTUBE_API_KEY (or paste a key) for an API fallback."))

    def _youtube_api_search(self, q: str, api_key: str) -> list[dict[str, Any]]:
        try:
            data = self._get_json(YT_API_SEARCH, {"part": "snippet", "q": q, "type": "video", "maxResults": 15, "key": api_key})
        except Exception:
            return []
        items = data.get("items", [])
        ids = list({i["snippet"]["channelId"] for i in items})
        handles: dict[str, str] = {}
        if ids:
            try:
                ch = self._get_json(YT_API_CHANNELS, {"part": "snippet", "id": ",".join(ids[:50]), "key": api_key})
                handles = {c["id"]: c["snippet"].get("customUrl", "") for c in ch.get("items", [])}
            except Exception:
                pass
        out = []
        for i in items:
            cid = i["snippet"]["channelId"]
            h = handles.get(cid, "")
            out.append({"id": i["id"].get("videoId", ""), "t": i["snippet"].get("title", ""),
                        "ch": f"https://www.youtube.com/{h}" if h.startswith("@") else f"https://www.youtube.com/channel/{cid}",
                        "ver": False})
        return out

    def _cbfc(self, titles, today, progress) -> None:
        targets = [t for t in titles if t.category == "Movie" and t.origin == "Indian"
                   and (t.from_list or not t.sort_date or t.sort_date >= (today - timedelta(days=120)).isoformat())]
        if not targets:
            return
        progress(88, "CBFC website: locating Search Film form")
        form = self._cbfc_form()
        if form is None:
            self.notes.append("CBFC's website did not serve its Search Film form to the server (it blocks bots); "
                              "certificates were not checked this run – use the skill's browser step for CBFC.")
            return
        action, method, fields, box = form
        found = 0
        for i, t in enumerate(targets[:CBFC_MAX_TITLES]):
            data = dict(fields)
            data[box] = t.title
            try:
                if method == "post":
                    r = self.http.session.post(action, data=data, timeout=self.http.timeout_seconds)
                else:
                    r = self.http.session.get(action, params=data, timeout=self.http.timeout_seconds)
                recs = parse_cbfc_results(r.text)
            except Exception:
                continue
            rec = next((x for x in recs if norm(x.get("film")) == norm(t.title)), None)
            if rec:
                found += 1
                t.certificate_type = t.certificate_type or rec.get("category", "")
                t.certificate_no = t.certificate_no or rec.get("certno", "")
                t.duration = t.duration or rec.get("length", "")
                t.cbfc_office = t.cbfc_office or rec.get("office", "")
                if t.status in ("Confirmed", "Tentative", "Date TBA"):
                    t.status = "Certified (CBFC)"
                t.source += f"; CBFC website, checked {today:%d %b %Y}"
            if i % 10 == 0:
                progress(88 + int(8 * i / len(targets)), f"CBFC website: {i}/{len(targets)} titles")
            time.sleep(0.3)
        self.notes.append(f"CBFC website: {found} of {min(len(targets), CBFC_MAX_TITLES)} recent/upcoming films matched a certificate.")

    def _cbfc_form(self):
        for url in CBFC_SEARCH_CANDIDATES + [CBFC_HOME]:
            try:
                html = self.http.get_text(url)
            except Exception:
                continue
            soup = BeautifulSoup(html, "html.parser")
            for form in soup.find_all("form"):
                box = form.find("input", attrs={"type": lambda v: v in (None, "text", "search")})
                if not box or not box.get("name"):
                    continue
                fields = {i.get("name"): i.get("value", "") for i in form.find_all("input") if i.get("name")}
                for s in form.find_all("select"):
                    if s.get("name"):
                        opt = s.find("option", selected=True) or s.find("option")
                        fields[s["name"]] = opt.get("value", "") if opt else ""
                action = form.get("action") or url
                if not action.startswith("http"):
                    action = re.sub(r"[^/]*$", "", url) + action.lstrip("/") if not action.startswith("/") else CBFC_HOME.rstrip("/") + action
                return action, (form.get("method") or "get").lower(), fields, box["name"]
        return None
